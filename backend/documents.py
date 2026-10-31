"""Document extraction and supported input types."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
import re
import unicodedata

from docx import Document
from pypdf import PdfReader


DOCUMENT_TYPES = {".txt", ".md", ".pdf", ".docx"}
CODE_LANGUAGES = {
    ".py": "python", ".js": "javascript", ".jsx": "javascript", ".ts": "typescript",
    ".tsx": "tsx", ".java": "java", ".c": "c", ".h": "c", ".cpp": "cpp",
    ".hpp": "cpp", ".cs": "c_sharp", ".go": "go", ".rs": "rust", ".php": "php",
    ".rb": "ruby", ".swift": "swift", ".kt": "kotlin", ".scala": "scala",
    ".html": "html", ".css": "css", ".sql": "sql", ".sh": "bash",
}
SUPPORTED_TYPES = DOCUMENT_TYPES | set(CODE_LANGUAGES)


def normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text).replace("\x00", " ")
    return re.sub(r"\s+", " ", text).strip()


def extract_content(filename: str, content: bytes, max_chars: int) -> tuple[str, str, str | None]:
    """Return extracted text, content kind and optional programming language."""
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_TYPES:
        raise ValueError(f"Unsupported file extension: {extension or '(none)'}")
    try:
        if extension in {".txt", ".md"} or extension in CODE_LANGUAGES:
            text = content.decode("utf-8-sig")
        elif extension == ".pdf":
            reader = PdfReader(BytesIO(content), strict=True)
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            doc = Document(BytesIO(content))
            chunks = [paragraph.text for paragraph in doc.paragraphs]
            chunks.extend(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
            text = "\n".join(chunks)
    except Exception as exc:
        raise ValueError(f"Could not read {extension} file: {exc}") from exc
    text = normalize_text(text)
    if not text:
        raise ValueError("No extractable text was found in the file")
    if len(text) > max_chars:
        raise ValueError(f"Extracted content exceeds the {max_chars} character limit")
    kind = "code" if extension in CODE_LANGUAGES else "document"
    return text, kind, CODE_LANGUAGES.get(extension)
